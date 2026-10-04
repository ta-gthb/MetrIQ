"""R 76-2 report fidelity: the document follows the template in force.

Audit item 10 asks for a report whose sections, tables and checklists are the
ones OIML R 76-2 prescribes, with MetrIQ's extra traceability as a layer on top
rather than as a substitute. The renderers therefore read the section structure
from the activated report template, and these tests hold them to it.
"""

from __future__ import annotations

import io
import uuid

import pytest

from app.database import SessionLocal
from app.models import EvaluationCase
from app.services.report_engine.docx import render_docx
from app.services.report_engine.pdf import render_pdf
from app.services.report_engine.sections import (
    heading_text,
    section_layer,
    section_number,
    section_required,
)
from app.services.report_engine.snapshot import build_report_snapshot

from tests.conftest import png_bytes as fixture_png
from tests.report_fixture import build_golden_case

TEMPLATE_SECTIONS = {
    "applicant": ("1", "Applicant and manufacturer"),
    "instrument": ("2", "Instrument identification and metrological characteristics"),
    "conditions": ("3", "Laboratory and test conditions"),
    "equipment": ("4", "Test equipment and traceability"),
    "summary": ("5", "Summary of applicable tests and results"),
    "tests": ("6", "Individual test records"),
    "checklist": ("7", "Construction and identification checklist"),
    "evidence": ("8", "Evidence and attachments"),
    "review": ("9", "Review, approval and signatures"),
    "versions": ("10", "Ruleset, template and revision history"),
    "verification": ("11", "Document verification"),
}


@pytest.fixture(scope="module")
def snapshot(client, tokens, observations) -> dict:
    """A fully recorded case, evidence included, rendered as a report snapshot."""

    db = SessionLocal()
    try:
        case_id = build_golden_case(client, tokens, db, observations)
    finally:
        db.close()
    # The instrument nameplate is the mandatory evidence; it is also the
    # photograph the generated report has to list.
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/attachments",
        files={"file": ("nameplate_photograph.png", fixture_png(), "image/png")},
        data={
            "category": "nameplate_photograph",
            "caption": "nameplate_photograph fixture",
            "auto_classify": "false",
        },
        headers=tokens["ENGINEER"],
    )
    assert uploaded.status_code == 201, uploaded.text

    db = SessionLocal()
    try:
        case = db.get(EvaluationCase, uuid.UUID(case_id))
        payload, _digest = build_report_snapshot(db, case=case, report_no="RPT-FIDELITY-1")
        return payload
    finally:
        db.close()


def _docx_text(payload: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _pdf_text(payload: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(payload)).pages)


def test_the_report_names_the_mode_it_was_written_in(snapshot):
    assert snapshot["versions"]["template_mode"] == "oiml_r76_2_aligned"
    assert snapshot["versions"]["template_mode_label"] == "OIML R 76-2 aligned"


def test_the_template_separates_the_standard_format_from_the_enhanced_layer(snapshot):
    """R 76-2 prescribes the report; the traceability block is MetrIQ's own and
    is declared as such instead of being presented as part of the standard."""
    assert section_layer(snapshot, "tests") == "oiml_r76_2"
    assert section_layer(snapshot, "review") == "oiml_r76_2"
    assert section_layer(snapshot, "verification") == "enhanced"


def test_every_section_is_titled_and_numbered_by_the_template(snapshot):
    for key, (number, title) in TEMPLATE_SECTIONS.items():
        assert section_number(snapshot, key) == number
        assert heading_text(snapshot, key, "fallback") == f"{number}. {title}"


def test_the_pdf_carries_every_section_the_template_prescribes(snapshot):
    text = _pdf_text(render_pdf(snapshot))
    for key, (number, title) in TEMPLATE_SECTIONS.items():
        if key == "cover":
            continue
        assert f"{number}. {title}" in text, f"the PDF is missing section {number} ({key})"


def test_the_docx_carries_every_section_the_template_prescribes(snapshot):
    text = _docx_text(render_docx(snapshot))
    for key, (number, title) in TEMPLATE_SECTIONS.items():
        if key == "cover":
            continue
        assert f"{number}. {title}" in text, f"the DOCX is missing section {number} ({key})"


def test_the_checklists_have_their_own_section_not_a_row_in_the_metrological_one(snapshot):
    """R 76-2 separates the construction and identification checklists from the
    metrological test records; the earlier renderers interleaved them."""
    checklist_tests = [item for item in snapshot["tests"] if item.get("layout") == "checklist"]
    assert checklist_tests, "the golden fixture must exercise the checklists"

    for text in (_pdf_text(render_pdf(snapshot)), _docx_text(render_docx(snapshot))):
        assert "7. Construction and identification checklist" in text
        for position, test in enumerate(checklist_tests, start=1):
            assert f"7.{position} {test['name']} ({test['test_code']})" in text
        # The metrological records keep their own numbering.
        assert "6.1 Weighing performance (T-WP)" in text


def test_a_template_without_a_section_map_falls_back_to_the_renderers_own_wording():
    """A missing or older template must not leave the report unnumbered."""
    bare = {"versions": {}, "meta": {}}
    assert heading_text(bare, "evidence", "Evidence and attachments") == "8. Evidence and attachments"
    assert section_number(bare, "evidence") == "8"
    assert section_number(bare, "not_a_section") == ""
    # An unknown key is treated as required, so a section can never vanish
    # because a template entry was mistyped.
    assert section_required(bare, "evidence") is True


def test_an_optional_section_with_nothing_to_report_is_left_out(snapshot):
    """A section the template marks optional disappears when it is empty, and a
    required one stays - so the reader never has to guess which it was."""
    import copy

    optional = copy.deepcopy(snapshot)
    optional["evidence"] = []
    for entry in optional["versions"]["template_sections"]:
        if entry["key"] == "evidence":
            entry["required"] = False
    assert section_required(optional, "evidence") is False

    text = _pdf_text(render_pdf(optional))
    assert "8. Evidence and attachments" not in text
    assert "No evidence attached." not in text

    # Marked required, the same empty section is rendered, empty and visible.
    required = copy.deepcopy(optional)
    for entry in required["versions"]["template_sections"]:
        if entry["key"] == "evidence":
            entry["required"] = True
    text = _pdf_text(render_pdf(required))
    assert "8. Evidence and attachments" in text
    assert "No evidence attached." in text